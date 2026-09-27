#!/usr/bin/env python3
"""Does the text on a slide actually fit, and does it collide with its neighbours?

THE BUG THIS EXISTS TO CATCH. On 26 September 2026 Adam reviewed session 10 on
paper and said the content was good and the formatting poor -- text too small,
and text overlapping text. Measuring his hand-edited copy against the build
found that he had raised 116 shapes to a consistent scale, and that **the build
he was sent already contained 30 overlapping text-box pairs**. Every mechanical
gate had passed it. `theme.py` enforces a 14 pt legibility floor and warns about
images that underfill their slot; nothing asked whether a string fits the box it
was put in, which is the failure a reader sees first and a script never sees.

HOW THE ESTIMATE WORKS, AND WHY IT IS ALLOWED TO BE ROUGH. Laying out text
exactly needs font metrics and a shaping engine. We do not need exact: we need
to know when a paragraph is obviously too long for its box. So a run of text is
modelled as

    characters per line = box width / (size_pt/72 * K)     K = 0.5
    height = lines * size_pt/72 * 1.21

K = 0.5 em is the average advance of proportional Latin text and was calibrated
against PowerPoint's own autofit heights in Adam's saved deck (318 boxes, mean
error -0.10 in, i.e. the estimate is slightly conservative -- it predicts a
little less height than PowerPoint allots, so a flagged box is really full).

Short strings also get a horizontal extent, so a slide number in the corner is
not reported as colliding with a full-width footer whose text stops well short
of it.

    python tools/slide_fit.py private/build/decks/Deck.pptx
"""
import math
import re
import sys
from pathlib import Path

NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
EMU = 914400.0
K = 0.5          # average glyph advance, in em
LINE = 1.21      # line height, in em
OVERFLOW_TOL = 0.10   # inches a box may be over before we care
COLLIDE_TOL = 0.06    # inches of overlap before we care


def _paragraphs(tf):
    """[(text, size_pt)] with the size each paragraph will actually render at."""
    out = []
    for p in tf.paragraphs:
        size = None
        for r in p.runs:
            rp = r._r.find(NS + "rPr")
            if rp is not None and rp.get("sz"):
                size = int(rp.get("sz")) / 100.0
                break
        if size is None and p._pPr is not None:
            d = p._pPr.find(NS + "defRPr")
            if d is not None and d.get("sz"):
                size = int(d.get("sz")) / 100.0
        if size is None:
            e = p._p.find(NS + "endParaRPr")
            if e is not None and e.get("sz"):
                size = int(e.get("sz")) / 100.0
        out.append(("".join(r.text for r in p.runs), size or 18.0))
    return out


def extent(paras, box_w):
    """(needed_height_in, widest_used_width_in) for these paragraphs."""
    h = 0.0
    used_w = 0.0
    for txt, size in paras:
        cw = size / 72.0 * K
        per = max(1, int(box_w / cw))
        n = max(1, math.ceil(len(txt) / per)) if txt else 1
        h += n * size / 72.0 * LINE
        used_w = max(used_w, min(box_w, len(txt) * cw))
    return h, used_w


def analyze(path):
    from pptx import Presentation
    pres = Presentation(str(path))
    overflow, collide, offslide = [], [], []
    SW, SH = pres.slide_width / EMU, pres.slide_height / EMU
    for si, s in enumerate(pres.slides, 1):
        live = []
        for sh in s.shapes:
            if not sh.has_text_frame:
                continue
            txt = re.sub(r"\s+", " ", sh.text_frame.text or "").strip()
            if not txt:
                continue
            if None in (sh.left, sh.top, sh.width, sh.height):
                continue
            l, t = sh.left / EMU, sh.top / EMU
            w, h = sh.width / EMU, sh.height / EMU
            if l < -0.02 or t < -0.02 or l + w > SW + 0.02 or t + h > SH + 0.02:
                offslide.append((si, txt[:44], round(l + w, 2), round(t + h, 2)))
            paras = _paragraphs(sh.text_frame)
            need, uw = extent(paras, w)
            if need > h + OVERFLOW_TOL:
                overflow.append((si, txt[:52], round(h, 2), round(need, 2),
                                 max(p[1] for p in paras)))
            live.append((txt, l, t, uw, max(h, need)))
        for i in range(len(live)):
            for j in range(i + 1, len(live)):
                a, b = live[i], live[j]
                ox = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
                oy = min(a[2] + a[4], b[2] + b[4]) - max(a[2], b[2])
                if ox > COLLIDE_TOL and oy > COLLIDE_TOL:
                    collide.append((si, round(ox, 2), round(oy, 2),
                                    a[0][:40], b[0][:40]))
    return overflow, collide, offslide


def report(path, indent="  "):
    """Print what does not fit. Returns the number of problems."""
    over, coll, off = analyze(path)
    if off:
        print(f"{indent}!! {len(off)} shape(s) run off the slide:")
        for si, txt, r, b in off[:10]:
            print(f"{indent}   s{si:<3} right {r}\" bottom {b}\"  | {txt}")
    if over:
        print(f"{indent}!! {len(over)} text box(es) overflow:")
        for si, txt, h, need, size in over[:14]:
            print(f"{indent}   s{si:<3} {size:>4.1f}pt  box {h}\" needs {need}\"  | {txt}")
        if len(over) > 14:
            print(f"{indent}   ... and {len(over) - 14} more")
    if coll:
        print(f"{indent}!! {len(coll)} overlapping text pair(s):")
        for si, ox, oy, a, b in coll[:14]:
            print(f'{indent}   s{si:<3} {ox}"x{oy}"  "{a}"  vs  "{b}"')
        if len(coll) > 14:
            print(f"{indent}   ... and {len(coll) - 14} more")
    return len(over) + len(coll) + len(off)


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    bad = 0
    for p in sys.argv[1:]:
        print(Path(p).name)
        bad += report(p)
    print(f"\n{bad} problem(s)." if bad else "\nEverything fits.")
    sys.exit(1 if bad else 0)
